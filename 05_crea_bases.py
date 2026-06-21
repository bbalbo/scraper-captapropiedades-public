import pandas as pd
import os
from datetime import datetime
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.styles import Alignment, Font
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import PatternFill
from openpyxl.styles.borders import Border, Side


ZONAS_PRIORIDAD = {
    'santiago-oriente': 350, 'santiago-centro': 700, 'santiago-sur': 600,
    'santiago-poniente': 500, 'gran-santiago': 2500,
    'nunoa': 130, 'las-condes': 150,
    'maipu': 350, 'la-florida': 350, 'puente-alto': 500, 'estacion-central': 350,
    'san-miguel': 95, 'chile-norte': 1500, 'costa-central': 1200,
    'chile-sur': 1500
}

ZONAS_GEOGRAFICAS = {
    "santiago-oriente": {"Las Condes", "Lo Barnechea", "Vitacura"},
    "santiago-centro": {"Santiago", "Recoleta", "Independencia", "Estación Central", "Providencia", "Ñuñoa", "Macul", "San Joaquín", "San Miguel"},
    "santiago-sur": {"La Pintana", "Puente Alto", "San Bernardo", "El Bosque", "La Granja", "Lo Espejo", "San Ramón"},
    "santiago-poniente": {"Pudahuel", "Cerrillos", "Renca", "Quinta Normal", "Lo Prado", "Cerro Navia", "Pedro Aguirre Cerda", "Quilicura"},
    "gran-santiago": {
        "Cerrillos", "Cerro Navia", "Colina", "Conchalí", "El Bosque", "Estación Central",
        "Huechuraba", "Independencia", "La Cisterna", "La Florida", "La Granja",
        "La Pintana", "La Reina", "Las Condes", "Lo Barnechea", "Lo Espejo",
        "Lo Prado", "Macul", "Maipú", "Ñuñoa", "Pedro Aguirre Cerda", "Peñalolén",
        "Providencia", "Pudahuel", "Puente Alto", "Quilicura", "Quinta Normal",
        "Recoleta", "Renca", "San Joaquín", "San Miguel", "San Ramón", "Santiago",
        "Vitacura", "Melipilla", "Padre Hurtado", "San Bernardo"
    },
    "nunoa": {"Ñuñoa"},
    "las-condes": {"Las Condes"},
    "maipu": {"Maipú"},
    "la-florida": {"La Florida"},
    "puente-alto": {"Puente Alto"},
    "estacion-central": {"Estación Central"},
    "san-miguel": {"San Miguel"},

    "chile-norte": {
        "Arica", "Iquique", "Alto Hospicio", "Antofagasta", "Calama", "Tocopilla",
        "Mejillones", "Copiapó", "Vallenar", "Huasco", "La Serena", "Coquimbo",
        "Ovalle", "Papudo"
    },

    "costa-central": {
        "Valparaíso", "Viña del Mar", "Concón", "Quintero", "Papudo", "Zapallar",
        "Algarrobo", "El Quisco", "El Tabo", "San Antonio", "Cartagena",
        "Santo Domingo", "Puchuncaví", "Casablanca", "Quillota", "Rancagua", "San Felipe"
    },

    "chile-sur": {
        "Concepción", "San Pedro de la Paz", "Talcahuano", "Hualpén", "Chiguayante", "Coronel",
        "Los Ángeles", "Chillán", "Chillán Viejo", "Linares", "Penco", "Tomé", "Temuco",
        "Padre las Casas", "Valdivia", "Osorno", "Puerto Varas", "Puerto Montt",
        "Frutillar", "Pucón", "Villarrica", "Coyhaique", "Punta Arenas"
    }
}


CARPETA_BASE = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'salida4')
SALIDA_ZONAS = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'bbdd_finales')
os.makedirs(SALIDA_ZONAS, exist_ok=True)

TIPOS_PROPIEDAD = {
    'casas': 'Casa',
    'departamentos': 'Departamento',
    'terrenos': 'Terreno'
}

HOY_STR = datetime.today().strftime("%Y-%m-%d")
df_total = []

for carpeta_nombre, nombre_propio in TIPOS_PROPIEDAD.items():
    carpeta_tipo = os.path.join(CARPETA_BASE, carpeta_nombre)
    if not os.path.isdir(carpeta_tipo):
        continue

    archivos_hoy = [f for f in os.listdir(carpeta_tipo)
                    if f.startswith(f"con_probabilidades_{HOY_STR}") and f.endswith(".xlsx")]
    for archivo in archivos_hoy:
        ruta = os.path.join(carpeta_tipo, archivo)
        df = pd.read_excel(ruta)

        if 'Ubicación' in df.columns:
            df.drop(columns=['Comuna'], errors='ignore', inplace=True)
            df.rename(columns={'Ubicación': 'Comuna'}, inplace=True)

        df.insert(0, 'Tipo de Propiedad', nombre_propio)
        df.columns = [col.replace('_', ' ') for col in df.columns]
        df_total.append(df)

if not df_total:
    print("⚠️ No se encontraron archivos de hoy")
    exit()

df = pd.concat(df_total, ignore_index=True)

# print("🧪 Columnas disponibles:", df.columns.tolist())

# LIMPIEZA Y PRIORIZACIÓN
df = df[df['Precio CLP'].fillna(0) > 1_000_000]
df = df.dropna(subset=['Comuna', 'Precio CLP', 'probabilidad corredor', 'Fecha Publicación'])
df['Fecha Publicación'] = pd.to_datetime(df['Fecha Publicación'], errors='coerce')

df = df.dropna(subset=['Fecha Publicación'])

for zona, comunas in ZONAS_GEOGRAFICAS.items():
    df_zona = df[
        (df['Comuna'].isin(comunas)) &
        (df['probabilidad corredor'] <= 0.35)
    ].copy()

    df_zona = df_zona.sort_values(by=['Fecha Publicación'], ascending=False)

    top_n = ZONAS_PRIORIDAD.get(zona, 100)
    df_zona = df_zona.head(top_n)

    if df_zona.empty:
        continue

    # === LIMPIEZA FINAL Y FORMATO DE PROBABILIDAD ===
    columnas_a_eliminar = ['Es Corredor (Heurística)', 'texto combinado', 'es corredor']
    df_zona.drop(columns=columnas_a_eliminar, inplace=True, errors='ignore')

    df_zona['probabilidad_num'] = df_zona['probabilidad corredor'] * 100
    df_zona['probabilidad corredor'] = df_zona['probabilidad_num'].map(lambda x: f"{x:.2f} %")

    # === Guardar el archivo limpio ===
    df_zona.drop(columns=['probabilidad_num'], inplace=True)
    nombre_archivo = os.path.join(SALIDA_ZONAS, f"{zona}.xlsx")
    df_zona.to_excel(nombre_archivo, index=False)

    # === Ajustes visuales con openpyxl ===
    wb = load_workbook(nombre_archivo)
    ws = wb.active
    ws.freeze_panes = "A2"
    # === FORMATO ENCABEZADOS ===
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="4F81BD")  # Azul suave
    header_alignment = Alignment(horizontal='center', vertical='center')

    thin_border = Border(
        left=Side(style='thin', color='DDDDDD'),
        right=Side(style='thin', color='DDDDDD'),
        top=Side(style='thin', color='DDDDDD'),
        bottom=Side(style='thin', color='DDDDDD')
    )

    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=ws.max_column):
        for cell in row:
            cell.border = thin_border

    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment

    # Opcional: altura de la fila de encabezado
    ws.row_dimensions[1].height = 24

    for idx, cell in enumerate(ws[1], start=1):
        col_letter = get_column_letter(idx)
        col_name = str(cell.value).strip().lower()
        col_cells = list(ws.iter_cols(min_col=idx, max_col=idx, min_row=2, max_row=ws.max_row))[0]

        if col_name == 'url':
            continue

        if col_name == 'descripción':
            ws.column_dimensions[col_letter].width = 30
            continue

        max_len = max([len(str(c.value)) if c.value else 0 for c in col_cells] + [len(col_name)])
        if col_name in ['título', 'precio clp']:
            ws.column_dimensions[col_letter].width = min(max_len + 2, 80) + 10
        else:
            ws.column_dimensions[col_letter].width = min(max_len + 2, 60)

        for c in col_cells:
            c.alignment = Alignment(horizontal='center')

        if col_name == 'precio clp':
            for c in col_cells:
                if c.value is not None:
                    c.number_format = u'[$$-es-CL]#,##0'
        elif col_name == 'precio uf':
            for c in col_cells:
                if c.value is not None:
                    c.number_format = '0'

        if col_name == 'probabilidad corredor':
            for c in col_cells:
                try:
                    value_float = float(str(c.value).replace('%', '').replace(',', '.'))
                    if value_float <= 20:
                        c.font = Font(color="7CFC00", bold=True, size=13)
                    elif value_float <= 30:
                        c.font = Font(color="E6E600", bold=True, size=13)
                    else:
                        c.font = Font(color="FF0000", bold=True, size=13)
                except:
                    pass
                c.alignment = Alignment(horizontal='center', vertical='center')

    wb.save(nombre_archivo)
    print(f"✅ Exportada zona {zona} con {len(df_zona)} registros")
